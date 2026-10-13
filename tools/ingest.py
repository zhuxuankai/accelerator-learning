# -*- coding: utf-8 -*-
"""增量入库：抽取 -> 分块 -> 写入 SQLite FTS5 索引

只读取 kb/config.json 里 include_dirs 指定的目录。
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    KB,
    MANIFEST_PATH,
    WS,
    chunk_text,
    connect_db,
    load_config,
    now_iso,
    read_json,
    sha1_file,
    slugify,
    tokenize,
    version_key,
    write_json,
)
from ocr_pdf import ocr_pages  # noqa: E402

TEXT_SUFFIXES = {".pdf", ".pptx", ".docx", ".md", ".txt", ".tex", ".csv", ".json"}
SKIP_SUFFIXES = {".exe", ".dll", ".png", ".jpg", ".jpeg", ".mat", ".zip", ".rar", ".7z", ".pyc"}


# ---------------------------------------------------------------- 抽取
def extract_pdf(path: Path) -> list[dict]:
    """先用 pypdf 逐页抽文本（快），稀疏页再用 pdfplumber 复核（准）。

    只有两种抽取器都拿不到内容的页才留给 OCR。
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    units = []
    sparse_pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        units.append({"unit": i, "text": text})
        if len(text.strip()) < 40:
            sparse_pages.append(i)

    if not sparse_pages:
        return units

    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for page_no in sparse_pages:
                if page_no > len(pdf.pages):
                    continue
                try:
                    alt = pdf.pages[page_no - 1].extract_text() or ""
                except Exception:
                    alt = ""
                if len(alt.strip()) > len(units[page_no - 1]["text"].strip()):
                    units[page_no - 1]["text"] = alt
                    units[page_no - 1]["extractor"] = "pdfplumber"
    except Exception as exc:
        print(f"[warn] pdfplumber 复核失败 {path.name}: {exc}")
    return units


def extract_pptx(path: Path) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(str(path))
    units = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text
                if txt and txt.strip():
                    parts.append(txt.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            parts.append(f"[备注] {notes}")
        units.append({"unit": i, "text": "\n".join(parts)})
    return units


def extract_docx(path: Path) -> list[dict]:
    import docx

    doc = docx.Document(str(path))
    buf: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            buf.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                buf.append(" | ".join(cells))
    return [{"unit": 1, "text": "\n".join(buf)}]


def extract_text_file(path: Path) -> list[dict]:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return [{"unit": 1, "text": path.read_text(encoding=enc)}]
        except (UnicodeDecodeError, LookupError):
            continue
    return [{"unit": 1, "text": ""}]


def extract(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".docx":
        return extract_docx(path)
    return extract_text_file(path)


def classify(units: list[dict], cfg: dict) -> tuple[str, list[int]]:
    """返回 (mode, 需要 OCR 的页码)。mode: none / sparse / full"""
    threshold = int(cfg.get("ocr", {}).get("sparse_char_threshold", 40))
    full_ratio = float(cfg.get("ocr", {}).get("full_image_ratio", 0.8))
    if not units:
        return "none", []
    sparse = [u["unit"] for u in units if len((u.get("text") or "").strip()) < threshold]
    ratio = len(sparse) / len(units)
    if ratio >= full_ratio:
        return "full", [u["unit"] for u in units]
    return "sparse", sparse


# ---------------------------------------------------------------- 主流程
def discover(cfg: dict) -> list[Path]:
    files: list[Path] = []
    for folder in cfg["include_dirs"]:
        base = WS / folder
        if not base.exists():
            continue
        for p in sorted(base.rglob("*")):
            if not p.is_file():
                continue
            if p.suffix.lower() in SKIP_SUFFIXES:
                continue
            if p.suffix.lower() not in TEXT_SUFFIXES:
                continue
            if p.name.startswith("~$"):
                continue
            files.append(p)
    return files


def doc_kind(rel_path: str) -> str:
    p = Path(rel_path)
    top = p.parts[0]
    if top == "文献":
        return "paper"
    if p.suffix.lower() == ".pptx":
        return "slide"
    if p.suffix.lower() == ".pdf":
        if any(k in p.stem for k in ("教科书", "教材", "Linear Accelerators")):
            return "book"
        return "slide"
    return "doc"


def index_doc(conn, rel_path: str, path: Path, cfg: dict, skip_ocr: bool = False) -> dict:
    t0 = time.time()
    units = extract(path)
    mode, ocr_units = classify(units, cfg)
    ocr_cfg = cfg.get("ocr", {})
    ocr_text: dict[int, str] = {}
    confidences: list[float] = []

    if ocr_units and path.suffix.lower() == ".pdf" and not skip_ocr:
        # 只 OCR 需要补的页；full 模式下整份 OCR
        ocr_text = ocr_pages(
            rel_path,
            ocr_units,
            int(ocr_cfg.get("dpi", 220)),
            int(ocr_cfg.get("workers", 3)),
            int(ocr_cfg.get("batch", 24)),
        )
        from ocr_pdf import load_cache

        cache = load_cache(rel_path)
        confidences = [cache[p].get("confidence", 0) for p in ocr_units if p in cache]

    merged: list[dict] = []
    ocr_used = 0
    for u in units:
        page = u["unit"]
        text = (u.get("text") or "").strip()
        is_ocr = 0
        ocr_len = len(ocr_text.get(page, "").strip())
        # 原生文本够长就保留原生文本；只有原生缺失或明显更短才采用 OCR
        if ocr_len and (len(text) < 40 or ocr_len > 2 * len(text)):
            text = ocr_text[page].strip()
            is_ocr = 1
            ocr_used += 1
        merged.append({"unit": page, "text": text, "ocr": is_ocr})

    doc_meta = {
        "rel_path": rel_path,
        "title": Path(rel_path).stem,
        "course": Path(rel_path).parts[0],
        "kind": doc_kind(rel_path),
        "sha1": sha1_file(path),
        "mtime": path.stat().st_mtime,
        "n_units": len(merged),
        "ocr_units": ocr_used,
        "ocr_ratio": (ocr_used / len(merged)) if merged else 0.0,
        "status": "ok",
        "note": "" if mode != "full" else "整份扫描版，已 OCR",
        "version_key": version_key(rel_path),
        "indexed_at": now_iso(),
    }

    # 写抽取结果
    out = KB / "extracted" / f"{slugify(rel_path)}.jsonl"
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8") as fh:
        for item in merged:
            fh.write(json.dumps(item, ensure_ascii=False) + "\n")

    # 写库
    conn.execute("DELETE FROM chunks WHERE rel_path = ?", (rel_path,))
    conn.execute("DELETE FROM docs WHERE rel_path = ?", (rel_path,))
    cur = conn.execute(
        """INSERT INTO docs(rel_path,title,course,kind,sha1,mtime,n_units,n_chunks,
                            ocr_units,ocr_ratio,status,note,version_key,version_rank,is_latest,indexed_at)
           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            doc_meta["rel_path"], doc_meta["title"], doc_meta["course"], doc_meta["kind"],
            doc_meta["sha1"], doc_meta["mtime"], doc_meta["n_units"], 0,
            doc_meta["ocr_units"], doc_meta["ocr_ratio"], doc_meta["status"], doc_meta["note"],
            doc_meta["version_key"], 0, 1, doc_meta["indexed_at"],
        ),
    )
    doc_id = cur.lastrowid

    max_chars = int(cfg.get("index", {}).get("max_chars", 800))
    overlap = int(cfg.get("index", {}).get("overlap", 100))
    kind = doc_meta["kind"]
    label = "页" if kind in ("slide", "book", "paper", "doc") else "张"
    if kind == "slide" and path.suffix.lower() == ".pptx":
        label = "张"

    n_chunks = 0
    for item in merged:
        pieces = chunk_text(item["text"], max_chars, overlap)
        for part, piece in enumerate(pieces):
            cur = conn.execute(
                """INSERT INTO chunks(doc_id,unit_label,part,text,n_chars,ocr,is_latest,course,kind,rel_path)
                   VALUES(?,?,?,?,?,?,?,?,?,?)""",
                (
                    doc_id,
                    f"{label}{item['unit']}",
                    part,
                    piece,
                    len(piece),
                    item["ocr"],
                    1,
                    doc_meta["course"],
                    kind,
                    rel_path,
                ),
            )
            conn.execute(
                "INSERT INTO chunks_fts(rowid, tok) VALUES(?,?)",
                (cur.lastrowid, tokenize(f"{doc_meta['title']} {piece}")),
            )
            n_chunks += 1

    conn.execute("UPDATE docs SET n_chunks = ? WHERE id = ?", (n_chunks, doc_id))
    doc_meta["n_chunks"] = n_chunks
    doc_meta["elapsed"] = round(time.time() - t0, 1)
    if confidences:
        doc_meta["ocr_confidence"] = round(sum(confidences) / len(confidences), 1)
    return doc_meta


def refresh_versions(conn) -> None:
    """同一 version_key 里 mtime 最新者为最新版，其余标记版本位次。"""
    rows = conn.execute("SELECT id, version_key, mtime FROM docs").fetchall()
    groups: dict[str, list] = {}
    for row in rows:
        groups.setdefault(row["version_key"], []).append(row)
    for key, items in groups.items():
        items.sort(key=lambda r: r["mtime"], reverse=True)
        for rank, row in enumerate(items, 1):
            is_latest = 1 if rank == 1 else 0
            conn.execute(
                "UPDATE docs SET version_rank = ?, is_latest = ? WHERE id = ?",
                (rank, is_latest, row["id"]),
            )
            conn.execute("UPDATE chunks SET is_latest = ? WHERE doc_id = ?", (is_latest, row["id"]))


def main() -> None:
    cfg = load_config()
    ap = argparse.ArgumentParser()
    ap.add_argument("--full", action="store_true", help="忽略清单，强制重新抽取全部文件")
    ap.add_argument("--only", help="只处理指定相对路径")
    ap.add_argument("--no-ocr", action="store_true", help="跳过 OCR（只做文本抽取）")
    args = ap.parse_args()

    if args.no_ocr:
        cfg.setdefault("ocr", {})

    manifest = read_json(MANIFEST_PATH, {"version": 1, "files": {}})
    known: dict = manifest.get("files", {})
    paths = discover(cfg)
    if args.only:
        paths = [p for p in paths if str(p.relative_to(WS)).replace("\\", "/") == args.only.replace("\\", "/")]

    added, changed, skipped, ocr_pending = [], [], [], []
    conn = connect_db()
    try:
        seen = set()
        for path in paths:
            rel = str(path.relative_to(WS)).replace("\\", "/")
            seen.add(rel)
            info = known.get(rel)
            digest = sha1_file(path)
            if info and info.get("sha1") == digest and info.get("status") == "ok" and not args.full:
                skipped.append(rel)
                continue
            (added if info is None else changed).append(rel)
            try:
                meta = index_doc(conn, rel, path, cfg, skip_ocr=args.no_ocr)
            except Exception as exc:  # 单文件失败不中断整批
                conn.rollback()
                known[rel] = {
                    "sha1": digest, "mtime": path.stat().st_mtime, "status": "error",
                    "note": str(exc)[:300], "last_indexed": now_iso(),
                }
                manifest["files"] = known
                write_json(MANIFEST_PATH, manifest)
                print(f"[error] {rel}: {exc}")
                continue
            known[rel] = {
                "sha1": meta["sha1"], "mtime": meta["mtime"], "status": meta["status"],
                "n_units": meta["n_units"], "n_chunks": meta["n_chunks"],
                "ocr_units": meta["ocr_units"], "note": meta["note"],
                "last_indexed": meta["indexed_at"],
            }
            print(
                f"[ok] {rel} 单元={meta['n_units']} 块={meta['n_chunks']} "
                f"OCR页={meta['ocr_units']} 用时={meta['elapsed']}s"
            )

        removed = [rel for rel in known if rel not in seen]
        for rel in removed:
            conn.execute("DELETE FROM chunks WHERE rel_path = ?", (rel,))
            conn.execute("DELETE FROM docs WHERE rel_path = ?", (rel,))
            known.pop(rel, None)

        refresh_versions(conn)
        conn.commit()
        manifest["files"] = known
        manifest["last_run"] = now_iso()
        write_json(MANIFEST_PATH, manifest)
    finally:
        conn.close()

    total_docs = connect_db(readonly=True).execute("SELECT COUNT(*) c FROM docs").fetchone()["c"]
    print(
        f"\n完成：新增 {len(added)}，更新 {len(changed)}，跳过 {len(skipped)}，"
        f"移除 {len(removed)}，索引内文档总数 {total_docs}"
    )
    if ocr_pending:
        print("待 OCR：", ", ".join(ocr_pending))


if __name__ == "__main__":
    main()
